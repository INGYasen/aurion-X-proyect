#!/usr/bin/env python3
"""Excel del examen modelo U2: despeje a mano (fórmulas) + evidencia Octave completa."""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
OCTAVE = ROOT / "octave"
EVID = ROOT / "evidencias"
OUT = ROOT / "excel" / "examen_modelo_u2_completo.xlsx"

NAVY = "1F4E79"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_YELLOW = PatternFill("solid", fgColor="FFF2CC")
FILL_GREEN = PatternFill("solid", fgColor="E2EFDA")
FILL_ORANGE = PatternFill("solid", fgColor="FCE4D6")
FILL_BLUE = PatternFill("solid", fgColor="DDEBF7")
FILL_PINK = PatternFill("solid", fgColor="F8CBAD")
FONT_WHITE = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
FONT_TITLE = Font(bold=True, size=14, color=NAVY, name="Calibri")
FONT_SEC = Font(bold=True, size=12, color=NAVY, name="Calibri")
FONT = Font(size=10, name="Calibri")
FONT_B = Font(bold=True, size=10, name="Calibri")
FONT_MONO = Font(size=9, name="Consolas")
WRAP = Alignment(wrap_text=True, vertical="center")
NUM = "0.0000000000"
SCI = "0.00E+00"

EXCEL_ERR = re.compile(r"^#(DIV/0!|VALUE!|REF!|NAME\?|NUM!|NULL!|N/A)$")


def det3(a, b, c, d, e, f, g, h, i) -> str:
    """Determinante 3×3 expandido (fórmula escalar, se puede copiar)."""
    return (
        f"({a}*({e}*{i}-{f}*{h})"
        f"-{b}*({d}*{i}-{f}*{g})"
        f"+{c}*({d}*{h}-{e}*{g}))"
    )


def add_check(ws, checks: list[tuple[str, str]], row: int, col: int, label: str) -> None:
    checks.append((label, f"{get_column_letter(col)}{row}"))


def verify_workbook(path: Path) -> None:
    """Falla si alguna fórmula evalúa a error de Excel (#DIV/0!, etc.)."""
    import formulas

    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    bad: list[str] = []
    for key, val in sol.items():
        if val.value is None:
            continue
        try:
            items = val.value.flat if hasattr(val.value, "flat") else [val.value]
        except Exception:
            items = [val.value]
        for item in items:
            s = str(item).strip()
            if EXCEL_ERR.match(s):
                bad.append(f"{key} -> {s}")
    if bad:
        raise RuntimeError("Fórmulas con error:\n" + "\n".join(bad[:30]))


def build_validacion(ws, checks_by_sheet: dict[str, list[tuple[str, str]]]) -> None:
    widths(ws, {1: 22, 2: 12, 3: 38, 4: 14})
    r = banner(ws, 1, "VALIDACIÓN DE INTEGRACIÓN (fórmulas)", FILL_NAVY, FONT_WHITE)
    r = prose(
        ws,
        r,
        "Cada fila lee una celda de prueba del ejercicio. Si aparece la palabra Revisar o un error de Excel, la hoja marca FALTA. "
        "Al abrir el archivo, Excel recalcula todas las fórmulas enlazadas.",
        FILL_YELLOW,
    )
    label_row(ws, r, ["Hoja", "Celda", "Prueba", "Estado"])
    r += 1
    first = r
    for sheet, items in checks_by_sheet.items():
        for label, addr in items:
            paint(ws.cell(r, 1, sheet))
            paint(ws.cell(r, 2, addr))
            paint(ws.cell(r, 3, label))
            ref = f"'{sheet}'!{addr}"
            paint(
                ws.cell(
                    r,
                    4,
                    f'=IF(ISERROR({ref}),"FALTA",IF(ISNUMBER(SEARCH("Revisar",{ref})),"FALTA","OK"))',
                ),
                bold=True,
                fill=FILL_GREEN,
            )
            r += 1
    last = r - 1
    r += 1
    ws.cell(r, 1, "ESTADO GLOBAL DEL LIBRO").font = FONT_TITLE
    paint(
        ws.cell(r, 2, f'=IF(COUNTIF(D{first}:D{last},"FALTA")>0,"REVISAR HOJAS","INTEGRADO SIN FALLAS")'),
        bold=True,
        fill=FILL_GREEN,
    )


def widths(ws, spec: dict[int, float]) -> None:
    for col, w in spec.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "Métodos numéricos — examen modelo U2"
    ws.oddFooter.right.text = "Hoja &P"


def banner(ws, r: int, text: str, fill, font) -> int:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    cell = ws.cell(r, 1, text)
    cell.font = font
    cell.alignment = WRAP
    for c in range(1, 9):
        ws.cell(r, c).fill = fill
        ws.cell(r, c).alignment = WRAP
    ws.row_dimensions[r].height = 22
    return r + 1


def prose(ws, r: int, text: str, fill=None) -> int:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    cell = ws.cell(r, 1, text)
    cell.font = FONT
    cell.alignment = WRAP
    lines = text.count("\n") + 1 + len(text) // 110
    ws.row_dimensions[r].height = min(72, 15 * max(1, lines))
    if fill:
        for c in range(1, 9):
            ws.cell(r, c).fill = fill
    return r + 1


def label_row(ws, r: int, labels: list[str]) -> None:
    for j, h in enumerate(labels, 1):
        cell = ws.cell(r, j, h)
        cell.fill = FILL_NAVY
        cell.font = FONT_WHITE
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[r].height = 22


def paint(cell, fmt=None, fill=None, bold=False) -> None:
    cell.font = FONT_B if bold else FONT
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill


def embed(ws, path: Path, anchor: str, w=640, h=360) -> None:
    if not path.is_file():
        ws[anchor] = f"(falta imagen {path.name})"
        return
    img = XLImage(str(path))
    img.width = w
    img.height = h
    ws.add_image(img, anchor)


def paste_log(ws, r: int, path: Path) -> int:
    r = banner(ws, r, "EVIDENCIA OCTAVE — salida completa (texto para copiar al informe)", FILL_GREEN, FONT_SEC)
    if not path.is_file():
        return prose(ws, r, "No se encontró el log.")
    for line in path.read_text(encoding="utf-8").splitlines():
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        cell = ws.cell(r, 1, line if line else " ")
        cell.font = FONT_MONO
        ws.row_dimensions[r].height = 14
        r += 1
    return r + 1


def header(ws, titulo: str, enunciado: list[str]) -> int:
    r = banner(ws, 1, "FACULTAD DE INGENIERÍA Y ARQUITECTURA", FILL_NAVY, FONT_WHITE)
    r = banner(ws, r, "EXAMEN DE UNIDAD DE MÉTODOS NUMÉRICOS  ·  Mg. Jessica Perez Rivera", FILL_NAVY, FONT_WHITE)
    r = prose(ws, r, "Nombre: ______________________________________________")
    r = banner(ws, r, titulo, FILL_BLUE, FONT_TITLE)
    r = banner(ws, r, "ENUNCIADO (transcrito del PDF del examen modelo)", FILL_BLUE, FONT_B)
    for line in enunciado:
        r = prose(ws, r, line)
    return r + 1


def build_indice(ws, checks_by_sheet: dict[str, list[tuple[str, str]]]) -> None:
    widths(ws, {1: 28, 2: 42, 3: 55, 4: 36, 5: 22, 6: 18, 7: 18, 8: 18})
    r = banner(ws, 1, "FACULTAD DE INGENIERÍA Y ARQUITECTURA — Métodos numéricos", FILL_NAVY, FONT_WHITE)
    r = banner(ws, r, "Examen modelo U2 — una hoja por ejercicio", FILL_NAVY, FONT_WHITE)
    r = prose(ws, r, "Nombre: ______________________________________________")
    r = prose(
        ws,
        r,
        "En cada hoja: (1) el enunciado, (2) el despeje a mano con fórmulas de Excel para poder copiar cómo salió cada número, "
        "(3) las iteraciones que pide el examen en aritmética exacta, (4) la corrida completa de Octave con gráfica incrustada y el log, "
        "(5) una prueba Excel contra Octave.",
        FILL_YELLOW,
    )
    r += 1
    label_row(ws, r, ["Hoja", "Ejercicio del PDF", "Qué se despeja a mano", "Evidencia Octave"])
    rows = [
        ("Ej1_No_lineal", "1. Sistema no lineal", "Jacobiana, 2 iteraciones de Newton en fracciones, despeje de punto fijo, sustitución de (±1,0,0) y soluciones reales", "Varios puntos de partida + 2 gráficas"),
        ("Ej2_Jacobi_Seidel", "2. Jacobi y Gauss-Seidel", "Reordenamiento E.D.D., despeje de x,y,z, dos iteraciones con fracciones", "Hasta error 1e-10, radio espectral, 2 gráficas"),
        ("Ej3_Newton_interp", "3. Catenaria, polinomio de Newton", "Tabla de diferencias divididas y evaluación anidada en x=200", "Curva, flecha y tabla"),
        ("Ej4_Minimos", "4. Pitot, parábola por mínimos cuadrados", "Sumas término a término, sistema normal y Cramer", "Parábola, v(4) y residuos"),
        ("Ej5_Lagrange", "5. Lagrange inverso, péndulo", "Cada factor L_i(y) escrito y evaluado en y=2", "Curva x=P(y) y el valor pedido"),
    ]
    r += 1
    for row in rows:
        for j, val in enumerate(row, 1):
            paint(ws.cell(r, j, val))
        ws.row_dimensions[r].height = 36
        for c in range(1, 5):
            ws.cell(r, c).alignment = WRAP
        r += 1
    r += 1
    r = prose(ws, r, "Regenerar evidencias y este archivo:", FILL_GREEN)
    r = prose(ws, r, "bash octave/ejecutar_todos.sh && python3 scripts/generar_examen_u2.py")
    r += 1
    n_checks = sum(len(v) for v in checks_by_sheet.values())
    ws.cell(r, 1, "Control de fórmulas:").font = FONT_B
    ws.cell(r, 2, f"Hoja 00_Validacion ({n_checks} pruebas enlazadas)")


def newton_step(ws, r: int, title: str, xb, yb, zb) -> dict[str, int]:
    """Un paso de Newton con J, F, det y Cramer en fórmulas escalares."""
    r = banner(ws, r, title, FILL_YELLOW, FONT_SEC)
    ws.cell(r, 1, "x , y , z").font = FONT_B
    paint(ws.cell(r, 2, xb), NUM, FILL_BLUE, True)
    paint(ws.cell(r, 3, yb), NUM, FILL_BLUE, True)
    paint(ws.cell(r, 4, zb), NUM, FILL_BLUE, True)
    rx = r
    r += 1
    ws.cell(r, 1, "F1 F2 F3").font = FONT_B
    paint(ws.cell(r, 2, f"=B{rx}^2+C{rx}^2-D{rx}"), NUM)
    paint(ws.cell(r, 3, f"=B{rx}^2+C{rx}^2+D{rx}^2-2"), NUM)
    paint(ws.cell(r, 4, f"=3*C{rx}+D{rx}-1"), NUM)
    rf = r
    r += 1
    ws.cell(r, 1, "J fila 1").font = FONT_B
    paint(ws.cell(r, 2, f"=2*B{rx}"), NUM)
    paint(ws.cell(r, 3, f"=2*C{rx}"), NUM)
    paint(ws.cell(r, 4, -1), NUM)
    j1 = r
    r += 1
    ws.cell(r, 1, "J fila 2").font = FONT_B
    paint(ws.cell(r, 2, f"=2*B{rx}"), NUM)
    paint(ws.cell(r, 3, f"=2*C{rx}"), NUM)
    paint(ws.cell(r, 4, f"=2*D{rx}"), NUM)
    j2 = r
    r += 1
    ws.cell(r, 1, "J fila 3").font = FONT_B
    paint(ws.cell(r, 2, 0), NUM)
    paint(ws.cell(r, 3, 3), NUM)
    paint(ws.cell(r, 4, 1), NUM)
    j3 = r
    r += 1
    ws.cell(r, 1, "det(J) expandido").font = FONT_B
    paint(
        ws.cell(
            r,
            2,
            "="
            + det3(
                f"B{j1}", f"C{j1}", f"D{j1}",
                f"B{j2}", f"C{j2}", f"D{j2}",
                f"B{j3}", f"C{j3}", f"D{j3}",
            ),
        ),
        "0.000000",
    )
    rd = r
    r += 1
    ws.cell(r, 1, "δx δy δz (Cramer)").font = FONT_B
    dx = det3(f"B{rf}", f"C{j1}", f"D{j1}", f"C{rf}", f"C{j2}", f"D{j2}", f"D{rf}", f"C{j3}", f"D{j3}")
    dy = det3(f"B{j1}", f"B{rf}", f"D{j1}", f"B{j2}", f"C{rf}", f"D{j2}", f"B{j3}", f"D{rf}", f"D{j3}")
    dz = det3(f"B{j1}", f"C{j1}", f"B{rf}", f"B{j2}", f"C{j2}", f"C{rf}", f"B{j3}", f"C{j3}", f"D{rf}")
    paint(ws.cell(r, 2, f"=({dx})/B{rd}"), NUM)
    paint(ws.cell(r, 3, f"=({dy})/B{rd}"), NUM)
    paint(ws.cell(r, 4, f"=({dz})/B{rd}"), NUM)
    rs = r
    r += 1
    ws.cell(r, 1, "X nuevo = X − δ").font = FONT_B
    paint(ws.cell(r, 2, f"=B{rx}-B{rs}"), NUM, FILL_ORANGE, True)
    paint(ws.cell(r, 3, f"=C{rx}-C{rs}"), NUM, FILL_ORANGE, True)
    paint(ws.cell(r, 4, f"=D{rx}-D{rs}"), NUM, FILL_ORANGE, True)
    return {"x": rx, "new": r, "next": r + 1}


def build_ej1(ws, checks: list[tuple[str, str]]) -> None:
    widths(ws, {i: 18 for i in range(1, 9)})
    ws.column_dimensions["A"].width = 28
    r = header(
        ws,
        "EJERCICIO 1 — Sistema de ecuaciones no lineales",
        [
            "x² + y² − z = 0",
            "x² + y² + z² = 2",
            "3y + z = 1",
            "a) Newton-Raphson y punto fijo: dos iteraciones desde (1, 0, 0).",
            "b) Mostrar si (±1, 0, 0) son soluciones del sistema escrito.",
            "c) Programa (Octave), varios puntos de partida.",
        ],
    )
    r = banner(ws, r, "A) DESARROLLO PARA COPIAR — Newton-Raphson (aritmética exacta)", FILL_YELLOW, FONT_SEC)
    for line in [
        "F1 = x² + y² − z ,   F2 = x² + y² + z² − 2 ,   F3 = 3y + z − 1.",
        "Jacobiana:  fila1 (2x, 2y, −1) ;  fila2 (2x, 2y, 2z) ;  fila3 (0, 3, 1).   Se resuelve J·δ = F y X ← X − δ.",
        "Iteración 0 en (1, 0, 0).  F = (1, −1, −1).",
        "J = [[2, 0, −1], [2, 0, 0], [0, 3, 1]].   det(J) = −1·(2·3 − 0·0) = −6.",
        "De la 2.ª ecuación: 2·δx = −1 ⇒ δx = −1/2.",
        "De la 1.ª: 2·(−1/2) − δz = 1 ⇒ −1 − δz = 1 ⇒ δz = −2.",
        "De la 3.ª: 3·δy − 2 = −1 ⇒ 3·δy = 1 ⇒ δy = 1/3.",
        "X1 = (1, 0, 0) − (−1/2, 1/3, −2) = (3/2, −1/3, 2).",
        "Iteración 1.  F(3/2, −1/3, 2) = (13/36, 157/36, 0).",
        "J = [[3, −2/3, −1], [3, −2/3, 4], [0, 3, 1]].  det(J) = 3·(−2/3 − 12) − (−2/3)·3 + (−1)·9 = −38 − 7 = −45.",
        "Cramer da δ = (59/180, −4/15, 4/5).",
        "X2 = (3/2 − 59/180, −1/3 + 4/15, 2 − 4/5) = (211/180, −1/15, 6/5).",
        "Las celdas de abajo recalculan lo mismo con fórmulas (para ver de dónde sale cada número) y deben coincidir con estas fracciones.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)
    r += 1
    step0 = newton_step(ws, r, "Comprobación con fórmulas — paso 0 → 1   (partida 1, 0, 0)", 1, 0, 0)
    step1 = newton_step(
        ws,
        step0["next"],
        "Comprobación con fórmulas — paso 1 → 2",
        f"=B{step0['new']}",
        f"=C{step0['new']}",
        f"=D{step0['new']}",
    )
    r = step1["next"]
    r = banner(ws, r, "Prueba del desarrollo a mano contra las fórmulas", FILL_ORANGE, FONT_SEC)
    ws.cell(r, 1, "X2 exacto (211/180, −1/15, 6/5)").font = FONT_B
    paint(ws.cell(r, 2, "=211/180"), NUM)
    paint(ws.cell(r, 3, "=-1/15"), NUM)
    paint(ws.cell(r, 4, "=6/5"), NUM)
    rex = r
    r += 1
    ws.cell(r, 1, "|X2 fórmulas − X2 exacto|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{step1['new']}-B{rex})"), SCI)
    paint(ws.cell(r, 3, f"=ABS(C{step1['new']}-C{rex})"), SCI)
    paint(ws.cell(r, 4, f"=ABS(D{step1['new']}-D{rex})"), SCI)
    r += 1
    ws.cell(r, 1, "Estado Newton a mano").font = FONT_B
    paint(ws.cell(r, 2, f'=IF(AND(B{r-1}<1E-9,C{r-1}<1E-9,D{r-1}<1E-9),"OK","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 2, "Newton: fórmulas = fracciones exactas")
    r += 2

    r = banner(ws, r, "A) DESARROLLO PARA COPIAR — Punto fijo (hay que despejar cada incógnita)", FILL_YELLOW, FONT_SEC)
    for line in [
        "De la ecuación 2, rama positiva (x0 = 1 > 0):   x = √(2 − y² − z²).",
        "De la ecuación 3:   y = (1 − z) / 3.",
        "De la ecuación 1:   z = x² + y².",
        "Actualización simultánea (solo se usan valores de la iteración anterior).",
        "Desde (1, 0, 0):   x1 = √2 ,   y1 = 1/3 ,   z1 = 1.",
        "Desde (√2, 1/3, 1):   x2 = √(2 − 1/9 − 1) = √(8/9) = 2√2 / 3 ,   y2 = 0 ,   z2 = 2 + 1/9 = 19/9.",
        "Tercera iteración: el argumento 2 − y2² − z2² = 2 − (19/9)² = −199/81 < 0. No hay siguiente raíz real con este despeje.",
        "El examen pide dos iteraciones: con este despeje existen y quedan escritas arriba. La tercera ya no es real; eso también se reporta.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)
    r += 1
    ws.cell(r, 1, "k").font = FONT_B
    label_row(ws, r, ["k", "x", "y", "z", "argumento 2−y²−z²", "¿siguiente real?"])
    r += 1
    # k=0 given
    paint(ws.cell(r, 1, 0))
    paint(ws.cell(r, 2, 1), NUM, FILL_BLUE)
    paint(ws.cell(r, 3, 0), NUM, FILL_BLUE)
    paint(ws.cell(r, 4, 0), NUM, FILL_BLUE)
    paint(ws.cell(r, 5, f"=2-C{r}^2-D{r}^2"), NUM)
    paint(ws.cell(r, 6, f'=IF(E{r}<0,"NO","SÍ")'))
    r0 = r
    r += 1
    for k in (1, 2, 3):
        prev = r - 1
        paint(ws.cell(r, 1, k))
        paint(ws.cell(r, 2, f'=IF(E{prev}<0,"—",SQRT(MAX(0,E{prev})))'), NUM)
        paint(ws.cell(r, 3, f'=IF(E{prev}<0,"—",(1-D{prev})/3)'), NUM)
        paint(ws.cell(r, 4, f'=IF(E{prev}<0,"—",B{prev}^2+C{prev}^2)'), NUM)
        paint(ws.cell(r, 5, f'=IF(NOT(ISNUMBER(C{r})),"—",2-C{r}^2-D{r}^2)'), NUM)
        paint(ws.cell(r, 6, f'=IF(E{prev}<0,"se detiene",IF(N(E{r})<0,"NO","SÍ"))'))
        r += 1
    r = prose(
        ws,
        r,
        "Prueba: x1 debe ser √2 ≈ 1.4142135624, y1 = 0.3333333333, z1 = 1; x2 = 2√2/3 ≈ 0.9428090416, y2 = 0, z2 = 19/9 ≈ 2.1111111111.",
        FILL_ORANGE,
    )
    r += 1

    r = banner(ws, r, "B) SUSTITUCIÓN PEDIDA — ¿(±1, 0, 0) anulan el sistema?", FILL_YELLOW, FONT_SEC)
    r = prose(
        ws,
        r,
        "El enunciado pide mostrar que (±1, 0, 0) son soluciones. Se sustituye en las tres ecuaciones. Si algún residuo no es 0, el punto no es solución del sistema escrito.",
        FILL_YELLOW,
    )
    label_row(ws, r, ["Punto", "F1 = x²+y²−z", "F2 = x²+y²+z²−2", "F3 = 3y+z−1", "¿Solución?"])
    r += 1
    for name, x, y, z in (("(1, 0, 0)", 1, 0, 0), ("(−1, 0, 0)", -1, 0, 0)):
        paint(ws.cell(r, 1, name), fill=FILL_BLUE)
        paint(ws.cell(r, 2, f"=({x})^2+({y})^2-({z})"), NUM)
        paint(ws.cell(r, 3, f"=({x})^2+({y})^2+({z})^2-2"), NUM)
        paint(ws.cell(r, 4, f"=3*({y})+({z})-1"), NUM)
        paint(ws.cell(r, 5, f'=IF(AND(ABS(B{r})<1E-12,ABS(C{r})<1E-12,ABS(D{r})<1E-12),"SÍ","NO")'), bold=True)
        r += 1
    for line in [
        "Conclusión de la sustitución: (±1, 0, 0) NO son soluciones. F1 = 1, F2 = −1, F3 = −1.",
        "DESPEJE ALGEBRAICO de las soluciones reales (esto sí se puede copiar como resolución exacta):",
        "Restar ec.1 de ec.2: (x²+y²+z²) − (x²+y²−z) = 2 ⇒ z² + z = 2 ⇒ z² + z − 2 = 0 ⇒ (z+2)(z−1) = 0.",
        "Caso z = 1: de ec.3, 3y+1 = 1 ⇒ y = 0. De ec.1, x² − 1 = 0 ⇒ x = ±1. Ec.2: 1+0+1 = 2, se cumple.",
        "Caso z = −2: 3y − 2 = 1 ⇒ y = 1. Ec.1: x² + 1 − (−2) = 0 ⇒ x² = −3, sin solución real.",
        "Soluciones reales del sistema: (1, 0, 1) y (−1, 0, 1). Abajo se comprueban con fórmulas.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)
    label_row(ws, r, ["Punto real", "F1", "F2", "F3", "¿Solución?"])
    r += 1
    for name, x, y, z in (("(1, 0, 1)", 1, 0, 1), ("(−1, 0, 1)", -1, 0, 1)):
        paint(ws.cell(r, 1, name), fill=FILL_GREEN)
        paint(ws.cell(r, 2, f"=({x})^2+({y})^2-({z})"), NUM)
        paint(ws.cell(r, 3, f"=({x})^2+({y})^2+({z})^2-2"), NUM)
        paint(ws.cell(r, 4, f"=3*({y})+({z})-1"), NUM)
        paint(ws.cell(r, 5, f'=IF(AND(ABS(B{r})<1E-12,ABS(C{r})<1E-12,ABS(D{r})<1E-12),"SÍ","NO")'), fill=FILL_GREEN, bold=True)
        r += 1
    r += 1

    ev = EVID / "ejercicio_01_newton_puntofijo"
    r = banner(ws, r, "C) EVIDENCIA OCTAVE — Newton desde varios puntos de partida (corrida completa)", FILL_GREEN, FONT_SEC)
    r = prose(ws, r, "Cada bloque es un punto inicial. No se recorta la tabla: se listan todas las iteraciones hasta ||F|| < 1e−12.")
    text = (ev / "resultado.txt").read_text(encoding="utf-8") if (ev / "resultado.txt").is_file() else ""
    # summary table of FIN lines and full ITER for first start only in a compact table, PLUS full log later
    label_row(ws, r, ["Arranque x", "Arranque y", "Arranque z", "Iteraciones", "x final", "y final", "z final", "||F|| final"])
    r += 1
    blocks = text.split("INICIO,")[1:]
    for block in blocks:
        lines = block.splitlines()
        x0, y0, z0 = lines[0].split(",")[:3]
        fin = [ln for ln in lines if ln.startswith("FIN,")]
        if not fin:
            continue
        _, ok, k, xf, yf, zf, nrm = fin[0].split(",")
        vals = [float(x0), float(y0), float(z0), int(k), float(xf), float(yf), float(zf), float(nrm)]
        for j, val in enumerate(vals, 1):
            paint(ws.cell(r, j, val), NUM if j != 4 else "0")
        r += 1
    r += 1
    r = prose(ws, r, "Detalle de TODAS las iteraciones del arranque pedido (1, 0, 0), para contrastar con las dos iteraciones a mano:", FILL_GREEN)
    label_row(ws, r, ["k", "x", "y", "z", "F1", "F2", "F3", "||F||"])
    r += 1
    first = blocks[0].splitlines() if blocks else []
    hand_row = None
    for ln in first:
        if not ln.startswith("ITER,"):
            continue
        parts = ln.split(",")
        k = int(parts[1])
        nums = [k] + [float(p) for p in parts[2:]]
        for j, val in enumerate(nums, 1):
            paint(ws.cell(r, j, val), "0" if j == 1 else (SCI if j >= 5 else NUM))
        if k == 2:
            hand_row = r
        r += 1
    if hand_row:
        r += 1
        ws.cell(r, 1, "Prueba |X2 a mano − X2 Octave|").font = FONT_B
        paint(ws.cell(r, 2, f"=ABS(B{step1['new']}-B{hand_row})"), SCI)
        paint(ws.cell(r, 3, f"=ABS(C{step1['new']}-C{hand_row})"), SCI)
        paint(ws.cell(r, 4, f"=ABS(D{step1['new']}-D{hand_row})"), SCI)
        r += 1
        paint(ws.cell(r, 2, f'=IF(AND(B{r-1}<1E-8,C{r-1}<1E-8,D{r-1}<1E-8),"OK contra Octave","Revisar")'), fill=FILL_GREEN, bold=True)
        add_check(ws, checks, r, 2, "Newton iter.2 Excel = Octave")
        r += 2
    r = banner(ws, r, "Prueba punto fijo — 2 iteraciones vs valores exactos", FILL_ORANGE, FONT_SEC)
    k1, k2 = r0 + 1, r0 + 2
    ws.cell(r, 1, "|x1 − √2| , |y1 − 1/3| , |z1 − 1|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{k1}-SQRT(2))"), SCI)
    paint(ws.cell(r, 3, f"=ABS(C{k1}-1/3)"), SCI)
    paint(ws.cell(r, 4, f"=ABS(D{k1}-1)"), SCI)
    r += 1
    ws.cell(r, 1, "|x2 − 2√2/3| , |y2| , |z2 − 19/9|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{k2}-2*SQRT(2)/3)"), SCI)
    paint(ws.cell(r, 3, f"=ABS(C{k2})"), SCI)
    paint(ws.cell(r, 4, f"=ABS(D{k2}-19/9)"), SCI)
    r += 1
    ws.cell(r, 1, "Estado punto fijo (2 iter.)").font = FONT_B
    paint(
        ws.cell(
            r,
            2,
            f'=IF(AND(B{r-1}<1E-9,C{r-1}<1E-9,D{r-1}<1E-9,B{r-2}<1E-9,C{r-2}<1E-9,D{r-2}<1E-9),"OK","Revisar")',
        ),
        fill=FILL_GREEN,
        bold=True,
    )
    add_check(ws, checks, r, 2, "Punto fijo: 2 iteraciones exactas")
    r += 2
    r = prose(ws, r, "Gráfica 1 — norma de F en escala logarítmica, todos los arranques (no se omiten los que ya son solución).", FILL_GREEN)
    embed(ws, ev / "octave_newton_varios.png", f"A{r+1}")
    r += 22
    r = prose(ws, r, "Gráfica 2 — las dos iteraciones reales del punto fijo desde (1, 0, 0).", FILL_GREEN)
    embed(ws, ev / "octave_punto_fijo.png", f"A{r+1}")
    r += 22
    paste_log(ws, r, ev / "resultado.txt")


def build_ej2(ws, checks: list[tuple[str, str]]) -> None:
    widths(ws, {i: 16 for i in range(1, 9)})
    ws.column_dimensions["A"].width = 34
    r = header(
        ws,
        "EJERCICIO 2 — Jacobi y Gauss-Seidel, aritmética exacta",
        [
            "Reorganizar el sistema para que sea estrictamente diagonalmente dominante (E.D.D.), o lo más parecido posible:",
            "−x + 4y − 2z = 7",
            "−y + 2z = 0",
            "−3x − z = 2",
            "Estudiar convergencia o divergencia de Jacobi y de Gauss-Seidel. Justificar. Calcular dos iteraciones. Aritmética exacta.",
            "El enunciado no da el vector inicial: se toma (0, 0, 0), que es el valor usual cuando no se indica otro.",
        ],
    )
    r = banner(ws, r, "DESARROLLO PARA COPIAR — por qué se reordenan las ecuaciones", FILL_YELLOW, FONT_SEC)
    for line in [
        "Criterio suficiente: |a_ii| > suma de |a_ij| con j distinto de i, en cada fila (E.D.D. estricta).",
        "Orden original (ec.1, ec.2, ec.3) para las incógnitas (x, y, z):",
        "Fila x: |−1| = 1  contra  |4| + |−2| = 6.   No.",
        "Fila y: |−1| = 1  contra  |2| = 2.   No.",
        "Fila z: |−1| = 1  contra  |−3| = 3.   No.",
        "La ecuación con mayor coeficiente en x es la tercera (|−3|). En y, la primera (|4|). En z, la segunda (|2|).",
        "Sistema reordenado:",
        "−3x − z = 2",
        "−x + 4y − 2z = 7",
        "−y + 2z = 0",
        "Fila x: 3 > 1.   Fila y: 4 > 1+2.   Fila z: 2 > 1.   Es E.D.D. estricta, así que Jacobi y Gauss-Seidel convergen.",
        "DESPEJE (las fórmulas de las celdas son este despeje):",
        "x = (−2 − z) / 3     porque −3x = 2 + z.",
        "y = (7 + x + 2z) / 4     porque 4y = 7 + x + 2z.",
        "z = y / 2     porque 2z = y.",
        "Jacobi usa solo la iteración anterior. Gauss-Seidel usa el x nuevo al calcular y, y el y nuevo al calcular z.",
        "Dos iteraciones desde (0,0,0), en fracciones:",
        "Jacobi 1: x = −2/3 , y = 7/4 , z = 0.",
        "Jacobi 2: x = −2/3 , y = (7 − 2/3)/4 = 19/12 , z = (7/4)/2 = 7/8.",
        "Gauss-Seidel 1: x = −2/3 , y = (7 − 2/3)/4 = 19/12 , z = (19/12)/2 = 19/24.",
        "Gauss-Seidel 2: x = (−2 − 19/24)/3 = −67/72 , y = (7 − 67/72 + 19/12)/4 = 551/288 , z = 551/576.",
        "Solución exacta del sistema reordenado (para la prueba): x = −1 , y = 2 , z = 1.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)

    r = banner(ws, r, "Comprobación del criterio E.D.D. con fórmulas (matriz reordenada)", FILL_YELLOW, FONT_SEC)
    label_row(ws, r, ["Ecuación", "a_x", "a_y", "a_z", "b", "|a_ii|", "suma |fuera|", "¿EDD?"])
    r += 1
    # rows: eq for x, y, z
    coefs = [(-3, 0, -1, 2), (-1, 4, -2, 7), (0, -1, 2, 0)]
    diag_col = (2, 3, 4)  # column of diagonal in the sheet (B=2 is ax)
    first_coef = r
    for i, (ax, ay, az, b) in enumerate(coefs):
        paint(ws.cell(r, 1, f"fila {i+1}"))
        paint(ws.cell(r, 2, ax), "0.00")
        paint(ws.cell(r, 3, ay), "0.00")
        paint(ws.cell(r, 4, az), "0.00")
        paint(ws.cell(r, 5, b), "0.00")
        # diagonal is column 2+i
        dcol = get_column_letter(2 + i)
        others = [get_column_letter(c) for c in (2, 3, 4) if c != 2 + i]
        paint(ws.cell(r, 6, f"=ABS({dcol}{r})"), "0.00")
        paint(ws.cell(r, 7, f"=ABS({others[0]}{r})+ABS({others[1]}{r})"), "0.00")
        paint(ws.cell(r, 8, f'=IF(F{r}>G{r},"SÍ","NO")'), bold=True)
        r += 1
    r += 1

    def iter_block(r: int, title: str, seidel: bool) -> tuple[int, int]:
        r = banner(ws, r, title, FILL_YELLOW, FONT_SEC)
        label_row(ws, r, ["k", "x", "y", "z", "fracción x", "fracción y", "fracción z", ""])
        r += 1
        paint(ws.cell(r, 1, 0))
        paint(ws.cell(r, 2, 0), NUM, FILL_BLUE)
        paint(ws.cell(r, 3, 0), NUM, FILL_BLUE)
        paint(ws.cell(r, 4, 0), NUM, FILL_BLUE)
        paint(ws.cell(r, 5, "0"))
        paint(ws.cell(r, 6, "0"))
        paint(ws.cell(r, 7, "0"))
        r += 1
        fracs = (
            [("-2/3", "7/4", "0"), ("-2/3", "19/12", "7/8")]
            if not seidel
            else [("-2/3", "19/12", "19/24"), ("-67/72", "551/288", "551/576")]
        )
        start = r
        for k in (1, 2):
            prev = r - 1
            paint(ws.cell(r, 1, k))
            # x always from previous z
            paint(ws.cell(r, 2, f"=(-2-D{prev})/3"), NUM, FILL_ORANGE, True)
            if seidel:
                paint(ws.cell(r, 3, f"=(7+B{r}+2*D{prev})/4"), NUM, FILL_ORANGE, True)
                paint(ws.cell(r, 4, f"=C{r}/2"), NUM, FILL_ORANGE, True)
            else:
                paint(ws.cell(r, 3, f"=(7+B{prev}+2*D{prev})/4"), NUM, FILL_ORANGE, True)
                paint(ws.cell(r, 4, f"=C{prev}/2"), NUM, FILL_ORANGE, True)
            fx, fy, fz = fracs[k - 1]
            paint(ws.cell(r, 5, fx))
            paint(ws.cell(r, 6, fy))
            paint(ws.cell(r, 7, fz))
            r += 1
        return r, start

    r, jstart = iter_block(r, "Jacobi — dos iteraciones (fórmulas = despeje; al lado, la fracción exacta)", False)
    r += 1
    ws.cell(r, 1, "Prueba Jacobi 2 vs fracciones").font = FONT_B
    paint(ws.cell(r, 2, f'=ABS(B{jstart+1}-(-2/3))'), SCI)
    paint(ws.cell(r, 3, f'=ABS(C{jstart+1}-(19/12))'), SCI)
    paint(ws.cell(r, 4, f'=ABS(D{jstart+1}-(7/8))'), SCI)
    r += 1
    paint(ws.cell(r, 2, f'=IF(AND(B{r-1}<1E-12,C{r-1}<1E-12,D{r-1}<1E-12),"OK","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 2, "Jacobi: 2 iter. = fracciones")
    r += 2
    r, gstart = iter_block(r, "Gauss-Seidel — dos iteraciones (usa el valor recién calculado)", True)
    r += 1
    ws.cell(r, 1, "Prueba GS 2 vs fracciones").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{gstart+1}-(-67/72))"), SCI)
    paint(ws.cell(r, 3, f"=ABS(C{gstart+1}-(551/288))"), SCI)
    paint(ws.cell(r, 4, f"=ABS(D{gstart+1}-(551/576))"), SCI)
    r += 1
    paint(ws.cell(r, 2, f'=IF(AND(B{r-1}<1E-12,C{r-1}<1E-12,D{r-1}<1E-12),"OK","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 2, "Gauss-Seidel: 2 iter. = fracciones")
    r += 2

    ev = EVID / "ejercicio_02_jacobi_seidel"
    text = (ev / "resultado.txt").read_text(encoding="utf-8")
    r = banner(ws, r, "EVIDENCIA OCTAVE — no solo dos iteraciones: hasta error < 1e−10, y el radio espectral", FILL_GREEN, FONT_SEC)
    rho_j = rho_g = None
    exact = (None, None, None)
    for ln in text.splitlines():
        if ln.startswith("RHO_JACOBI"):
            rho_j = float(ln.split(",")[1])
        elif ln.startswith("RHO_GS"):
            rho_g = float(ln.split(",")[1])
        elif ln.startswith("SOLUCION_EXACTA"):
            exact = tuple(float(x) for x in ln.split(",")[1:])
    ws.cell(r, 1, "ρ(B Jacobi)").font = FONT_B
    paint(ws.cell(r, 2, rho_j), NUM, FILL_GREEN, True)
    ws.cell(r, 3, "ρ(B Gauss-Seidel)").font = FONT_B
    paint(ws.cell(r, 4, rho_g), NUM, FILL_GREEN, True)
    r += 1
    ws.cell(r, 1, "¿ρ < 1? (converge)").font = FONT_B
    paint(ws.cell(r, 2, f'=IF(AND(B{r-1}<1,D{r-1}<1),"SÍ, ambos","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 2, "Convergencia: radios espectrales < 1")
    r += 1
    ws.cell(r, 1, "Solución exacta x, y, z").font = FONT_B
    for j, val in enumerate(exact, 2):
        paint(ws.cell(r, j, val), NUM, FILL_ORANGE, True)
    r += 2
    r = prose(ws, r, "Tabla completa de Jacobi (todas las iteraciones de la corrida):", FILL_GREEN)
    label_row(ws, r, ["k", "x", "y", "z", "error máx."])
    r += 1
    for ln in text.splitlines():
        if ln.startswith("JACOBI,"):
            _, k, x, y, z, err = ln.split(",")
            for j, val in enumerate((int(k), float(x), float(y), float(z), float(err)), 1):
                paint(ws.cell(r, j, val), "0" if j == 1 else (SCI if j == 5 else NUM))
            r += 1
    r += 1
    r = prose(ws, r, "Tabla completa de Gauss-Seidel:", FILL_GREEN)
    label_row(ws, r, ["k", "x", "y", "z", "error máx."])
    r += 1
    for ln in text.splitlines():
        if ln.startswith("GAUSS_SEIDEL,"):
            _, k, x, y, z, err = ln.split(",")
            for j, val in enumerate((int(k), float(x), float(y), float(z), float(err)), 1):
                paint(ws.cell(r, j, val), "0" if j == 1 else (SCI if j == 5 else NUM))
            r += 1
    r += 1
    r = prose(ws, r, "Gráfica — error de las dos corridas completas (Jacobi necesita más iteraciones; GS baja más rápido, ρ más chico).", FILL_GREEN)
    embed(ws, ev / "octave_convergencia.png", f"A{r+1}")
    r += 22
    r = prose(ws, r, "Gráfica — solución exacta frente al último vector de cada método.", FILL_GREEN)
    embed(ws, ev / "octave_solucion.png", f"A{r+1}")
    r += 22
    paste_log(ws, r, ev / "resultado.txt")


def build_ej3(ws, checks: list[tuple[str, str]]) -> None:
    widths(ws, {i: 18 for i in range(1, 9)})
    ws.column_dimensions["A"].width = 32
    r = header(
        ws,
        "EJERCICIO 3 — Polinomio de Newton y flecha de la catenaria",
        [
            "Torres de una línea de alta tensión, vano de 400 m. Alturas medidas del cable:",
            "Distancia (m):  0 , 120 , 230 , 310 , 400",
            "Altura (m):  50 , 43.28 , 42.18 , 44.42 , 50.00",
            "Obtener la forma de Newton del polinomio de interpolación y usarla para la flecha máxima en el punto medio de la catenaria.",
            "Punto medio del vano: x = 200 m. Las dos torres están a 50 m, así que la flecha es 50 − P(200).",
        ],
    )
    r = banner(ws, r, "DESARROLLO PARA COPIAR — diferencias divididas (cada celda es una resta dividida)", FILL_YELLOW, FONT_SEC)
    for line in [
        "f[x_i] = y_i.",
        "f[x_i, x_{i+1}] = (y_{i+1} − y_i) / (x_{i+1} − x_i).",
        "Órdenes siguientes: (diferencia de la columna anterior) / (x del extremo derecho − x del extremo izquierdo).",
        "A mano, con estos datos:",
        "f[x0,x1] = (43.28 − 50) / 120 = −6.72/120 = −0.056",
        "f[x1,x2] = (42.18 − 43.28) / 110 = −1.10/110 = −0.01",
        "f[x2,x3] = (44.42 − 42.18) / 80 = 2.24/80 = 0.028",
        "f[x3,x4] = (50 − 44.42) / 90 = 5.58/90 = 0.062",
        "f[x0,x1,x2] = (−0.01 − (−0.056)) / 230 = 0.046/230 = 0.0002",
        "f[x1,x2,x3] = (0.028 − (−0.01)) / 190 = 0.038/190 = 0.0002",
        "f[x2,x3,x4] = (0.062 − 0.028) / 170 = 0.034/170 = 0.0002",
        "El tercer orden sale 0 (0.0002 − 0.0002). El cuarto orden también es 0. El polinomio es de grado 2.",
        "P(x) = 50 − 0.056·(x − 0) + 0.0002·(x − 0)·(x − 120).",
        "En x = 200, forma anidada: se empieza por el coeficiente más alto (0).",
        "b2 = 0.0002",
        "b1 = −0.056 + (200 − 120)·0.0002 = −0.056 + 0.016 = −0.040",
        "b0 = 50 + (200 − 0)·(−0.040) = 50 − 8 = 42",
        "P(200) = 42 m.   Flecha máxima en el punto medio = 50 − 42 = 8 m.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)

    r = banner(ws, r, "Tabla de diferencias divididas (fórmulas). La primera fila de cada orden es el coeficiente de Newton.", FILL_YELLOW, FONT_SEC)
    label_row(ws, r, ["i", "x", "f[ ]", "1.er orden", "2.º orden", "3.er orden", "4.º orden"])
    r += 1
    data = [(0, 50), (120, 43.28), (230, 42.18), (310, 44.42), (400, 50)]
    rd = r
    for i, (x, y) in enumerate(data):
        paint(ws.cell(r, 1, i))
        paint(ws.cell(r, 2, x), "0", FILL_BLUE)
        paint(ws.cell(r, 3, y), "0.00", FILL_BLUE)
        r += 1
    # fill divided differences
    for i in range(4):
        rr = rd + i
        paint(ws.cell(rr, 4, f"=(C{rr+1}-C{rr})/(B{rr+1}-B{rr})"), "0.0000000000")
    for i in range(3):
        rr = rd + i
        paint(ws.cell(rr, 5, f"=(D{rr+1}-D{rr})/(B{rr+2}-B{rr})"), "0.0000000000")
    for i in range(2):
        rr = rd + i
        paint(ws.cell(rr, 6, f"=(E{rr+1}-E{rr})/(B{rr+3}-B{rr})"), SCI)
    paint(ws.cell(rd, 7, f"=(F{rd+1}-F{rd})/(B{rd+4}-B{rd})"), SCI)
    r = rd + 5 + 1
    r = banner(ws, r, "Evaluación anidada en x = 200 (cada línea es el paso que se copia)", FILL_YELLOW, FONT_SEC)
    ws.cell(r, 1, "x medio").font = FONT_B
    paint(ws.cell(r, 2, 200), "0", FILL_BLUE, True)
    rx = r
    r += 1
    # b4 = a4 = G_rd
    ws.cell(r, 1, "b4 = a4").font = FONT_B
    paint(ws.cell(r, 2, f"=G{rd}"), SCI)
    r += 1
    ws.cell(r, 1, "b3 = a3 + (x−x3)·b4").font = FONT_B
    paint(ws.cell(r, 2, f"=F{rd}+($B${rx}-B{rd+3})*B{r-1}"), SCI)
    r += 1
    ws.cell(r, 1, "b2 = a2 + (x−x2)·b3").font = FONT_B
    paint(ws.cell(r, 2, f"=E{rd}+($B${rx}-B{rd+2})*B{r-1}"), "0.0000000000")
    r += 1
    ws.cell(r, 1, "b1 = a1 + (x−x1)·b2").font = FONT_B
    paint(ws.cell(r, 2, f"=D{rd}+($B${rx}-B{rd+1})*B{r-1}"), "0.0000000000")
    r += 1
    ws.cell(r, 1, "P(200) = a0 + (x−x0)·b1").font = FONT_B
    paint(ws.cell(r, 2, f"=C{rd}+($B${rx}-B{rd})*B{r-1}"), "0.0000000000", FILL_ORANGE, True)
    rp = r
    r += 1
    ws.cell(r, 1, "Flecha = 50 − P(200)").font = FONT_B
    paint(ws.cell(r, 2, f"=50-B{rp}"), "0.0000000000", FILL_ORANGE, True)
    rf = r
    r += 1
    ws.cell(r, 1, "Prueba |P(200) − 42| y |flecha − 8|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{rp}-42)"), SCI)
    paint(ws.cell(r, 3, f"=ABS(B{rf}-8)"), SCI)
    paint(ws.cell(r, 4, f'=IF(AND(B{r}<1E-9,C{r}<1E-9),"OK","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 4, "Newton interp.: P(200) y flecha")
    r += 2

    ev = EVID / "ejercicio_03_newton_interp"
    text = (ev / "resultado.txt").read_text(encoding="utf-8")
    p200 = flecha = None
    for ln in text.splitlines():
        if ln.startswith("P200"):
            p200 = float(ln.split(",")[1])
        elif ln.startswith("FLECHA"):
            flecha = float(ln.split(",")[1])
    r = banner(ws, r, "EVIDENCIA OCTAVE", FILL_GREEN, FONT_SEC)
    ws.cell(r, 1, "P(200) Octave").font = FONT_B
    paint(ws.cell(r, 2, p200), NUM, FILL_GREEN, True)
    ws.cell(r, 3, "Flecha Octave").font = FONT_B
    paint(ws.cell(r, 4, flecha), NUM, FILL_GREEN, True)
    r += 1
    ws.cell(r, 1, "|Excel − Octave| en P(200)").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{rp}-B{r-1})"), SCI)
    paint(ws.cell(r, 3, f'=IF(B{r}<1E-8,"OK","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 3, "P(200) Excel = Octave")
    r += 2
    r = prose(ws, r, "Gráfica: datos, polinomio, cota de las torres (50 m) y el segmento de la flecha en x = 200.", FILL_GREEN)
    embed(ws, ev / "octave_catenaria.png", f"A{r+1}")
    r += 22
    paste_log(ws, r, ev / "resultado.txt")


def build_ej4(ws, checks: list[tuple[str, str]]) -> None:
    widths(ws, {i: 16 for i in range(1, 9)})
    ws.column_dimensions["A"].width = 28
    r = header(
        ws,
        "EJERCICIO 4 — Velocidad en un tubo (Pitot), parábola por mínimos cuadrados",
        [
            "Tubería circular de diámetro interior 20 cm. r es la distancia al centro, en cm.",
            "v (cm/s):  600 , 550 , 450 , 312 , 240",
            "r (cm):  0 , 3 , 5 , 7 , 8",
            "a) Curva v = f(r) que aproxima los datos.  b) Velocidad en r = 4 cm.",
            "Emplear ecuación parabólica por mínimos cuadrados: v = a·r² + b·r + c.",
        ],
    )
    r = banner(ws, r, "DESARROLLO PARA COPIAR — sumas y ecuaciones normales", FILL_YELLOW, FONT_SEC)
    for line in [
        "Con n puntos, el modelo v = a r² + b r + c lleva al sistema normal (incógnitas c, b, a):",
        "n·c + (Σr)·b + (Σr²)·a = Σv",
        "(Σr)·c + (Σr²)·b + (Σr³)·a = Σ(r v)",
        "(Σr²)·c + (Σr³)·b + (Σr⁴)·a = Σ(r² v)",
        "Cada producto está en una columna. La suma es la fórmula SUM de esa columna: así se ve de dónde sale cada coeficiente.",
        "Luego el sistema 3×3 se resuelve por Cramer. El determinante de cada celda está expandido (a11(a22·a33−a23·a32)−…).",
        "Al final, v(4) = a·16 + b·4 + c.",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)
    label_row(ws, r, ["i", "r", "v", "r²", "r³", "r⁴", "r·v", "r²·v"])
    r += 1
    datos = [(0, 600), (3, 550), (5, 450), (7, 312), (8, 240)]
    rd = r
    for i, (rv, vv) in enumerate(datos, 1):
        paint(ws.cell(r, 1, i))
        paint(ws.cell(r, 2, rv), "0", FILL_BLUE)
        paint(ws.cell(r, 3, vv), "0", FILL_BLUE)
        paint(ws.cell(r, 4, f"=B{r}^2"), "0")
        paint(ws.cell(r, 5, f"=B{r}^3"), "0")
        paint(ws.cell(r, 6, f"=B{r}^4"), "0")
        paint(ws.cell(r, 7, f"=B{r}*C{r}"), "0")
        paint(ws.cell(r, 8, f"=D{r}*C{r}"), "0")
        r += 1
    last = r - 1
    ws.cell(r, 1, "SUMAS").font = FONT_B
    paint(ws.cell(r, 2, f"=SUM(B{rd}:B{last})"), "0", FILL_YELLOW, True)  # Sr
    paint(ws.cell(r, 3, f"=SUM(C{rd}:C{last})"), "0", FILL_YELLOW, True)  # Sv
    paint(ws.cell(r, 4, f"=SUM(D{rd}:D{last})"), "0", FILL_YELLOW, True)  # Sr2
    paint(ws.cell(r, 5, f"=SUM(E{rd}:E{last})"), "0", FILL_YELLOW, True)  # Sr3
    paint(ws.cell(r, 6, f"=SUM(F{rd}:F{last})"), "0", FILL_YELLOW, True)  # Sr4
    paint(ws.cell(r, 7, f"=SUM(G{rd}:G{last})"), "0", FILL_YELLOW, True)  # Srv
    paint(ws.cell(r, 8, f"=SUM(H{rd}:H{last})"), "0", FILL_YELLOW, True)  # Sr2v
    rs = r
    n = len(datos)
    r += 2
    r = prose(ws, r, f"n = {n}.  Sistema, copiando las sumas de la fila {rs}:", FILL_YELLOW)
    # Matrix rows for unknowns c, b, a
    # [ n , Sr , Sr2 | Sv ]
    # [ Sr, Sr2, Sr3 | Srv ]
    # [ Sr2, Sr3, Sr4 | Sr2v ]
    label_row(ws, r, ["", "coef c", "coef b", "coef a", "término"])
    r += 1
    m0 = r
    paint(ws.cell(r, 1, "ec. 1"))
    paint(ws.cell(r, 2, n), "0")
    paint(ws.cell(r, 3, f"=B{rs}"), "0")
    paint(ws.cell(r, 4, f"=D{rs}"), "0")
    paint(ws.cell(r, 5, f"=C{rs}"), "0")
    r += 1
    paint(ws.cell(r, 1, "ec. 2"))
    paint(ws.cell(r, 2, f"=B{rs}"), "0")
    paint(ws.cell(r, 3, f"=D{rs}"), "0")
    paint(ws.cell(r, 4, f"=E{rs}"), "0")
    paint(ws.cell(r, 5, f"=G{rs}"), "0")
    r += 1
    paint(ws.cell(r, 1, "ec. 3"))
    paint(ws.cell(r, 2, f"=D{rs}"), "0")
    paint(ws.cell(r, 3, f"=E{rs}"), "0")
    paint(ws.cell(r, 4, f"=F{rs}"), "0")
    paint(ws.cell(r, 5, f"=H{rs}"), "0")
    m2 = r
    r += 1
    # det of columns B C D
    ws.cell(r, 1, "det(A) Cramer").font = FONT_B
    paint(ws.cell(r, 2, "=" + det3(f"B{m0}", f"C{m0}", f"D{m0}", f"B{m0+1}", f"C{m0+1}", f"D{m0+1}", f"B{m2}", f"C{m2}", f"D{m2}")), "0.00")
    rdet = r
    r += 1
    ws.cell(r, 1, "c = det(A_c)/det").font = FONT_B
    dc = det3(f"E{m0}", f"C{m0}", f"D{m0}", f"E{m0+1}", f"C{m0+1}", f"D{m0+1}", f"E{m2}", f"C{m2}", f"D{m2}")
    db = det3(f"B{m0}", f"E{m0}", f"D{m0}", f"B{m0+1}", f"E{m0+1}", f"D{m0+1}", f"B{m2}", f"E{m2}", f"D{m2}")
    da = det3(f"B{m0}", f"C{m0}", f"E{m0}", f"B{m0+1}", f"C{m0+1}", f"E{m0+1}", f"B{m2}", f"C{m2}", f"E{m2}")
    paint(ws.cell(r, 2, f"=({dc})/B{rdet}"), "0.0000000000", FILL_ORANGE, True)
    rc = r
    r += 1
    ws.cell(r, 1, "b = det(A_b)/det").font = FONT_B
    paint(ws.cell(r, 2, f"=({db})/B{rdet}"), "0.0000000000", FILL_ORANGE, True)
    rb = r
    r += 1
    ws.cell(r, 1, "a = det(A_a)/det").font = FONT_B
    paint(ws.cell(r, 2, f"=({da})/B{rdet}"), "0.0000000000", FILL_ORANGE, True)
    ra = r
    r += 1
    ws.cell(r, 1, "v(4) = a·16 + b·4 + c").font = FONT_B
    paint(ws.cell(r, 2, f"=B{ra}*16+B{rb}*4+B{rc}"), "0.0000000000", FILL_ORANGE, True)
    rv = r
    r += 2
    r = prose(ws, r, "Residuos punto a punto (v − v̂). Sirven para copiar la tabla del ajuste.", FILL_YELLOW)
    label_row(ws, r, ["r", "v medido", "v̂ = a r²+b r+c", "residuo"])
    r += 1
    first_res = r
    for i in range(5):
        src = rd + i
        paint(ws.cell(r, 1, f"=B{src}"), "0")
        paint(ws.cell(r, 2, f"=C{src}"), "0.00")
        paint(ws.cell(r, 3, f"=$B${ra}*A{r}^2+$B${rb}*A{r}+$B${rc}"), "0.0000")
        paint(ws.cell(r, 4, f"=B{r}-C{r}"), "0.0000")
        r += 1
    last_res = r - 1
    ws.cell(r, 1, "SSE = Σ residuo²").font = FONT_B
    paint(ws.cell(r, 2, f"=SUMPRODUCT(D{first_res}:D{last_res},D{first_res}:D{last_res})"), "0.0000")
    r += 2

    ev = EVID / "ejercicio_04_minimos"
    kv = {}
    for ln in (ev / "resultado.txt").read_text(encoding="utf-8").splitlines():
        if ln[:2] in ("a,", "b,", "c,") or ln.startswith("v4") or ln.startswith("SSE"):
            k, val = ln.split(",", 1)
            kv[k] = float(val)
    r = banner(ws, r, "EVIDENCIA OCTAVE y prueba contra el Cramer de esta hoja", FILL_GREEN, FONT_SEC)
    label_row(ws, r, ["", "Valor Octave", "Excel (Cramer)", "|dif.|", "Estado"])
    r += 1
    for label, row_excel, val_oct in (
        ("a", ra, kv["a"]),
        ("b", rb, kv["b"]),
        ("c", rc, kv["c"]),
        ("v(4)", rv, kv["v4"]),
    ):
        paint(ws.cell(r, 1, label))
        paint(ws.cell(r, 2, val_oct), "0.0000000000", FILL_GREEN)
        paint(ws.cell(r, 3, f"=B{row_excel}"), "0.0000000000", FILL_ORANGE)
        paint(ws.cell(r, 4, f"=ABS(C{r}-B{r})"), SCI)
        paint(ws.cell(r, 5, f'=IF(D{r}<1E-6,"OK","Revisar")'), bold=True, fill=FILL_GREEN)
        add_check(ws, checks, r, 5, f"Mínimos cuadrados: {label} Excel=Octave")
        r += 1
    r += 1
    r = prose(ws, r, "Gráfica del ajuste y del punto r = 4 cm.", FILL_GREEN)
    embed(ws, ev / "octave_parabola.png", f"A{r+1}")
    r += 22
    r = prose(ws, r, "Gráfica de residuos en cada r medido.", FILL_GREEN)
    embed(ws, ev / "octave_residuos.png", f"A{r+1}")
    r += 22
    paste_log(ws, r, ev / "resultado.txt")


def build_ej5(ws, checks: list[tuple[str, str]]) -> None:
    widths(ws, {i: 18 for i in range(1, 9)})
    ws.column_dimensions["A"].width = 42
    r = header(
        ws,
        "EJERCICIO 5 — Interpolación inversa con Lagrange de grado 2",
        [
            "y es la amplitud de un péndulo largo (cm) y x el tiempo (min).",
            "Puntos:  0  1  2  3  4  5  6",
            "x:  0.0  2.5  5.0  7.5  10.0  12.5  15.0",
            "y:  10.00  4.97  2.47  1.22  0.61  0.30  0.14",
            "Polinomio de Lagrange de segundo grado por los puntos (1), (2) y (3), y el x que corresponde a y = 2 cm.",
            "Interpolación inversa: se intercambian los papeles. Se construye x = P(y) y se evalúa en y = 2.",
            "Punto (1): y=4.97 , x=2.5.    Punto (2): y=2.47 , x=5.0.    Punto (3): y=1.22 , x=7.5.",
        ],
    )
    r = banner(ws, r, "DESARROLLO PARA COPIAR — cada factor del polinomio de Lagrange", FILL_YELLOW, FONT_SEC)
    for line in [
        "L1(y) = (y − y2)(y − y3) / ((y1 − y2)(y1 − y3))",
        "L2(y) = (y − y1)(y − y3) / ((y2 − y1)(y2 − y3))",
        "L3(y) = (y − y1)(y − y2) / ((y3 − y1)(y3 − y2))",
        "P(y) = x1·L1(y) + x2·L2(y) + x3·L3(y)",
        "Abajo el numerador y el denominador van en celdas distintas, para copiar la cuenta y no solo el resultado.",
        "Comprobación útil: L1+L2+L3 = 1 en cualquier y (partición de la unidad).",
    ]:
        r = prose(ws, r, line, FILL_YELLOW)
    ws.cell(r, 1, "y buscado").font = FONT_B
    paint(ws.cell(r, 2, 2), "0.00", FILL_BLUE, True)
    yq = r
    r += 1
    label_row(ws, r, ["", "punto (1)", "punto (2)", "punto (3)"])
    r += 1
    ws.cell(r, 1, "y_i").font = FONT_B
    for j, val in enumerate((4.97, 2.47, 1.22), 2):
        paint(ws.cell(r, j, val), "0.00", FILL_BLUE)
    ry = r
    r += 1
    ws.cell(r, 1, "x_i").font = FONT_B
    for j, val in enumerate((2.5, 5.0, 7.5), 2):
        paint(ws.cell(r, j, val), "0.00", FILL_BLUE)
    rx = r
    r += 2
    # numerators and denominators
    triples = [
        ("L1", 2, 3, 4),
        ("L2", 3, 2, 4),
        ("L3", 4, 2, 3),
    ]
    # columns: B is point1 (col 2), C point2 (col 3), D point3 (col 4)
    # For Li, own column, the other two
    specs = [
        ("L1 numerador (y−y2)(y−y3)", f"=($B${yq}-C{ry})*($B${yq}-D{ry})"),
        ("L1 denominador (y1−y2)(y1−y3)", f"=(B{ry}-C{ry})*(B{ry}-D{ry})"),
        ("L2 numerador (y−y1)(y−y3)", f"=($B${yq}-B{ry})*($B${yq}-D{ry})"),
        ("L2 denominador (y2−y1)(y2−y3)", f"=(C{ry}-B{ry})*(C{ry}-D{ry})"),
        ("L3 numerador (y−y1)(y−y2)", f"=($B${yq}-B{ry})*($B${yq}-C{ry})"),
        ("L3 denominador (y3−y1)(y3−y2)", f"=(D{ry}-B{ry})*(D{ry}-C{ry})"),
    ]
    rows_nd = []
    for title, formula in specs:
        ws.cell(r, 1, title).font = FONT_B
        paint(ws.cell(r, 2, formula), "0.000000")
        rows_nd.append(r)
        r += 1
    r += 1
    lrows = []
    for i, name in enumerate(("L1(2)", "L2(2)", "L3(2)")):
        num_r, den_r = rows_nd[2 * i], rows_nd[2 * i + 1]
        ws.cell(r, 1, f"{name} = numerador/denominador").font = FONT_B
        paint(ws.cell(r, 2, f"=B{num_r}/B{den_r}"), "0.000000000", FILL_ORANGE, True)
        lrows.append(r)
        r += 1
    ws.cell(r, 1, "Suma L1+L2+L3 (debe ser 1)").font = FONT_B
    paint(ws.cell(r, 2, f"=B{lrows[0]}+B{lrows[1]}+B{lrows[2]}"), "0.000000000")
    r += 1
    ws.cell(r, 1, "x(y=2) = x1 L1 + x2 L2 + x3 L3").font = FONT_B
    paint(
        ws.cell(r, 2, f"=B{rx}*B{lrows[0]}+C{rx}*B{lrows[1]}+D{rx}*B{lrows[2]}"),
        "0.000000000",
        FILL_ORANGE,
        True,
    )
    rans = r
    r += 1
    ws.cell(r, 1, "Prueba |suma Li − 1|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{r-2}-1)"), SCI)
    paint(ws.cell(r, 3, f'=IF(B{r}<1E-9,"OK partición","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 3, "Lagrange: L1+L2+L3 = 1")
    r += 2

    ev = EVID / "ejercicio_05_lagrange"
    kv = {}
    for ln in (ev / "resultado.txt").read_text(encoding="utf-8").splitlines():
        if "," in ln and ln.split(",")[0] in ("L1", "L2", "L3", "x_en_y2", "suma_L"):
            k, val = ln.split(",", 1)
            kv[k] = float(val)
    r = banner(ws, r, "EVIDENCIA OCTAVE", FILL_GREEN, FONT_SEC)
    ws.cell(r, 1, "x(2) Octave").font = FONT_B
    paint(ws.cell(r, 2, kv["x_en_y2"]), "0.000000000", FILL_GREEN, True)
    r += 1
    ws.cell(r, 1, "|Excel − Octave|").font = FONT_B
    paint(ws.cell(r, 2, f"=ABS(B{rans}-B{r-1})"), SCI)
    paint(ws.cell(r, 3, f'=IF(B{r}<1E-8,"OK Octave","Revisar")'), fill=FILL_GREEN, bold=True)
    add_check(ws, checks, r, 3, "Lagrange inverso: x(y=2) Excel=Octave")
    r += 2
    r = prose(ws, r, "Gráfica: x en función de y, los tres puntos usados y la lectura en y = 2 cm.", FILL_GREEN)
    embed(ws, ev / "octave_lagrange.png", f"A{r+1}")
    r += 22
    paste_log(ws, r, ev / "resultado.txt")


def main() -> None:
    subprocess.run(["bash", str(OCTAVE / "ejecutar_todos.sh")], check=True)
    wb = Workbook()
    wb.remove(wb.active)
    checks_by_sheet: dict[str, list[tuple[str, str]]] = {}
    ws_ind = wb.create_sheet("00_Indice", 0)
    ws_val = wb.create_sheet("00_Validacion", 1)
    sheets = [
        ("Ej1_No_lineal", build_ej1),
        ("Ej2_Jacobi_Seidel", build_ej2),
        ("Ej3_Newton_interp", build_ej3),
        ("Ej4_Minimos", build_ej4),
        ("Ej5_Lagrange", build_ej5),
    ]
    for name, fn in sheets:
        ws = wb.create_sheet(name)
        checks: list[tuple[str, str]] = []
        fn(ws, checks)
        checks_by_sheet[name] = checks
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110
        ws.oddHeader.left.text = "Métodos numéricos — examen modelo U2"
        ws.print_title_rows = "1:2"
        ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5, header=0.25, footer=0.25)
        ws.sheet_properties.tabColor = NAVY
    build_validacion(ws_val, checks_by_sheet)
    build_indice(ws_ind, checks_by_sheet)
    ws_ind.sheet_properties.tabColor = "C65911"
    ws_ind.sheet_view.showGridLines = False
    ws_val.sheet_properties.tabColor = "548235"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcOnSave = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    verify_workbook(OUT)
    print(f"Generado y verificado (0 errores de fórmula): {OUT}")


if __name__ == "__main__":
    main()
