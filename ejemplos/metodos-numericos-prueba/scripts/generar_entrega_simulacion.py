#!/usr/bin/env python3
"""Genera Excel multi-hoja: 1 hoja por ejercicio + imágenes Octave + pruebas."""
from __future__ import annotations

import os
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OCTAVE = ROOT / "octave"
EVID = ROOT / "evidencias"
EXCEL_OUT = ROOT / "excel" / "simulacion_examen_u2_multihoja.xlsx"

HEADER_LINES = [
    "FACULTAD DE INGENIERÍA Y ARQUITECTURA",
    "Simulación entrega U2 — Métodos Numéricos",
    "Una hoja por ejercicio | Evidencia Octave incrustada | Prueba Excel ↔ Octave",
]

thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE = Font(bold=True, size=13, color="1F4E79")
LBL = Font(bold=True, size=10)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
OK_FONT = Font(bold=True, color="006100")


def run_octave_scripts() -> None:
    script = OCTAVE / "ejecutar_todos.sh"
    subprocess.run(["bash", str(script)], check=True)


def sheet_header(ws, row_start: int = 1) -> int:
    r = row_start
    for line in HEADER_LINES:
        ws.cell(r, 1, line).font = TITLE if r == row_start else Font(size=10)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        r += 1
    return r + 1


def style_table_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def embed_image(ws, img_path: Path, anchor: str, width: int = 520, height: int = 300) -> None:
    if not img_path.is_file():
        ws[anchor.split("0")[0] if anchor[1].isdigit() else anchor] = "(sin imagen Octave)"
        return
    img = XLImage(str(img_path))
    img.width = width
    img.height = height
    ws.add_image(img, anchor)


def parse_kv(path: Path) -> dict[str, str]:
    data: dict[str, str] = {}
    if not path.is_file():
        return data
    for line in path.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.startswith(("i,", "row,", "xi,")):
            k, v = line.split("=", 1)
            if "," not in k:
                data[k.strip()] = v.strip()
    return data


def data_lines(path: Path) -> list[str]:
    return [ln for ln in path.read_text(encoding="utf-8").splitlines() if re.match(r"^\d+,", ln)]


def build_indice(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicios incluidos (simulación tipo examen modelo U2)").font = LBL
    r += 2
    hdr = ["#", "Tema", "Hoja Excel", "Detectado en PDF modelo"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, len(hdr))
    r += 1
    items = [
        ("1", "Bisección (raíz)", "Ejercicio_1", "—"),
        ("2", "Newton-Raphson", "Ejercicio_2", "—"),
        ("3", "Gauss 3×3", "Ejercicio_3", "—"),
        ("4", "Parábola mínimos cuadrados", "Ejercicio_4", "Sí (enunciado U2)"),
        ("5", "Integración (trapecio)", "Ejercicio_5", "Sí (enunciado U2)"),
    ]
    for row_data in items:
        for j, val in enumerate(row_data, 1):
            ws.cell(r, j, val).border = BORDER
        r += 1
    r += 2
    ws.cell(r, 1, "Regenerar:").font = LBL
    ws.cell(r, 2, "bash octave/ejecutar_todos.sh && python3 scripts/generar_entrega_simulacion.py")


def build_ej1(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicio 1 — Bisección").font = LBL
    r += 1
    ws.cell(r, 1, "f(x)=x³−x−2 en [1, 2], ε=10⁻⁶ (todo el ejercicio en esta hoja).")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    r += 2
    ev = EVID / "ejercicio_01_biseccion"
    kv = parse_kv(ev / "resultado.txt")
    raiz_o = float(kv.get("raiz", 0))

    ws.cell(r, 1, "Raíz Octave (referencia):").font = LBL
    ws.cell(r, 2, raiz_o).number_format = "0.0000000000"
    r_ref = r
    r += 2

    headers = ["i", "a", "b", "c", "f(c)", "error"]
    for j, h in enumerate(headers, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, len(headers))
    r0 = r + 1
    lines = data_lines(ev / "resultado.txt")
    show = lines if len(lines) <= 15 else lines[:10] + lines[-3:]
    for ln in show:
        parts = ln.split(",")
        for j, val in enumerate(parts[:6], 1):
            cell = ws.cell(r0, j, float(val) if j > 1 else int(val))
            cell.border = BORDER
            if j > 1:
                cell.number_format = "0.0000000000" if j < 6 else "0.00E+00"
        r0 += 1
    if len(lines) > 15:
        ws.cell(r0, 1, f"(Tabla: {len(lines)} filas en Octave; se muestran 10 + 3 finales)")
        r0 += 1
    r0 += 1
    last_data_row = r0 - 1
    ws.cell(r0, 1, "c final (última fila tabla):").font = LBL
    ws.cell(r0, 2, f"=D{last_data_row}")
    r0 += 1
    ws.cell(r0, 1, "Prueba |c_tabla − c_Octave|:").font = LBL
    ws.cell(r0, 2, f"=ABS(B{r0-1}-B{r_ref})")
    ws.cell(r0, 3, "Estado:")
    ws.cell(r0, 4, f'=IF(B{r0}<1E-8,"OK","Revisar")')
    r0 += 2
    ws.cell(r0, 1, "Evidencia Octave (imagen):").font = LBL
    embed_image(ws, ev / "octave_grafica.png", f"A{r0+1}")


def build_ej2(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicio 2 — Newton-Raphson").font = LBL
    r += 1
    ws.cell(r, 1, "f(x)=cos(x)−x, x₀=0.8, parar cuando |f(x)|<10⁻⁸.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    r += 2
    ev = EVID / "ejercicio_02_newton"
    kv = parse_kv(ev / "resultado.txt")
    raiz_o = float(kv.get("raiz", 0))
    ws.cell(r, 1, "Raíz Octave:").font = LBL
    ws.cell(r, 2, raiz_o).number_format = "0.0000000000"
    r_ref = r
    r += 2
    hdr = ["i", "x", "f(x)", "f'(x)", "error"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, len(hdr))
    r0 = r + 1
    for ln in data_lines(ev / "resultado.txt"):
        parts = ln.split(",")
        for j, val in enumerate(parts[:5], 1):
            cell = ws.cell(r0, j, float(val) if j > 1 else int(val))
            cell.border = BORDER
            if j > 1:
                cell.number_format = "0.00E+00" if j in (3, 4, 5) else "0.0000000000"
        r0 += 1
    r0 += 1
    ws.cell(r0, 1, "x final tabla:").font = LBL
    ws.cell(r0, 2, f"=B{r0-1}")
    r0 += 1
    ws.cell(r0, 1, "Prueba |x_tabla − x_Octave|:").font = LBL
    ws.cell(r0, 2, f"=ABS(B{r0-1}-B{r_ref})")
    ws.cell(r0, 3, "Estado:")
    ws.cell(r0, 4, f'=IF(B{r0}<1E-10,"OK","Revisar")')
    r0 += 2
    ws.cell(r0, 1, "Evidencia Octave:").font = LBL
    embed_image(ws, ev / "octave_grafica.png", f"A{r0+1}")


def build_ej3(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicio 3 — Eliminación gaussiana").font = LBL
    r += 1
    ws.cell(r, 1, "Sistema 3×3 (matriz en script Octave).")
    r += 2
    ev = EVID / "ejercicio_03_gauss"
    text = (ev / "resultado.txt").read_text(encoding="utf-8")
    m = re.search(r"sol,x1=([^,]+),x2=([^,]+),x3=(.+)", text)
    x1, x2, x3 = map(float, m.groups()) if m else (0.0, 0.0, 0.0)
    kv = parse_kv(ev / "resultado.txt")
    res = float(kv.get("residual_norm", 0))
    hdr = ["x₁", "x₂", "x₃"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, 3)
    r += 1
    for j, val in enumerate([x1, x2, x3], 1):
        c = ws.cell(r, j, val)
        c.number_format = "0.0000000000"
        c.border = BORDER
    r += 2
    ws.cell(r, 1, "||Ax−b|| (Octave):").font = LBL
    ws.cell(r, 2, res).number_format = "0.00E+00"
    ws.cell(r, 3, "Estado:")
    ws.cell(r, 4, "OK" if res < 1e-6 else "Revisar")
    r += 2
    ws.cell(r, 1, "Evidencia Octave:").font = LBL
    embed_image(ws, ev / "octave_grafica.png", f"A{r+1}")


def build_ej4(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicio 4 — Parábola por mínimos cuadrados").font = LBL
    r += 1
    ws.cell(
        r,
        1,
        "Emplear ecuación parabólica y = ax²+bx+c (como examen modelo U2).",
    )
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    r += 2
    ev = EVID / "ejercicio_04_minimos_cuadrados"
    kv = parse_kv(ev / "resultado.txt")
    a, b, c = float(kv["a"]), float(kv["b"]), float(kv["c"])
    ws.cell(r, 1, "Coeficientes Octave → celdas Excel:").font = LBL
    r += 1
    ws.cell(r, 1, "a"); ws.cell(r, 2, a).number_format = "0.0000000000"
    ws.cell(r, 3, "b"); ws.cell(r, 4, b).number_format = "0.0000000000"
    ws.cell(r, 5, "c"); ws.cell(r, 6, c).number_format = "0.0000000000"
    row_coef = r
    r += 2
    hdr = ["i", "xᵢ", "yᵢ", "ŷ Octave", "ŷ Excel", "Δ"]
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, len(hdr))
    r0 = r + 1
    for ln in data_lines(ev / "resultado.txt"):
        i, xi, yi, yhat = ln.split(",")
        ws.cell(r0, 1, int(i)).border = BORDER
        ws.cell(r0, 2, float(xi)).border = BORDER
        ws.cell(r0, 3, float(yi)).border = BORDER
        ws.cell(r0, 4, float(yhat)).border = BORDER
        ws.cell(r0, 5, f"=$B${row_coef}*B{r0}^2+$D${row_coef}*B{r0}+$F${row_coef}").border = BORDER
        ws.cell(r0, 5).number_format = "0.0000"
        ws.cell(r0, 6, f"=ABS(D{r0}-E{r0})").border = BORDER
        ws.cell(r0, 6).number_format = "0.00E+00"
        r0 += 1
    first, last = r + 1, r0 - 1
    r0 += 1
    ws.cell(r0, 1, "Prueba Σ(y−ŷ)² (datos vs Excel):").font = LBL
    ws.cell(r0, 2, f"=SUMXMY2(C{first}:C{last},E{first}:E{last})")
    ws.cell(r0, 3, "Estado:")
    ws.cell(r0, 4, f'=IF(B{r0}<1E-6,"OK","Revisar")')
    r0 += 2
    ws.cell(r0, 1, "Evidencia Octave (datos + parábola):").font = LBL
    embed_image(ws, ev / "octave_grafica.png", f"A{r0+1}")


def build_ej5(ws) -> None:
    r = sheet_header(ws)
    ws.cell(r, 1, "Ejercicio 5 — Integración (trapecio)").font = LBL
    r += 1
    ws.cell(r, 1, "∫₀¹ e^(−x²) dx — refinamiento de n (simulación U2).")
    r += 2
    ev = EVID / "ejercicio_05_trapecio"
    hdr = ["n", "h", "Integral Octave", "Error vs ref."]
    for j, h in enumerate(hdr, 1):
        ws.cell(r, j, h)
    style_table_header(ws, r, len(hdr))
    r0 = r + 1
    for ln in data_lines(ev / "resultado.txt"):
        parts = ln.split(",")
        for j, val in enumerate(parts[:4], 1):
            cell = ws.cell(r0, j, float(val) if j > 1 else int(val))
            cell.border = BORDER
            if j >= 3:
                cell.number_format = "0.0000000000" if j == 3 else "0.00E+00"
        r0 += 1
    kv = parse_kv(ev / "resultado.txt")
    ref = float(kv.get("referencia", 0))
    r0 += 1
    ws.cell(r0, 1, "Referencia Octave:").font = LBL
    ws.cell(r0, 2, ref).number_format = "0.0000000000"
    ws.cell(r0, 3, "Mejor trapecio (n=64):").font = LBL
    ws.cell(r0, 4, f"=C{r0-1}")
    r0 += 1
    ws.cell(r0, 1, "Prueba error final:").font = LBL
    ws.cell(r0, 2, f"=D{r0-1}")
    ws.cell(r0, 3, "Estado:")
    ws.cell(r0, 4, f'=IF(B{r0}<1E-5,"OK","Revisar")')
    r0 += 2
    ws.cell(r0, 1, "Evidencia Octave:").font = LBL
    embed_image(ws, ev / "octave_grafica.png", f"A{r0+1}")


def main() -> None:
    run_octave_scripts()
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Indice"
    build_indice(ws0)
    for title, fn in [
        ("Ejercicio_1", build_ej1),
        ("Ejercicio_2", build_ej2),
        ("Ejercicio_3", build_ej3),
        ("Ejercicio_4", build_ej4),
        ("Ejercicio_5", build_ej5),
    ]:
        ws = wb.create_sheet(title)
        fn(ws)
        for col in range(1, 11):
            ws.column_dimensions[get_column_letter(col)].width = 15
    EXCEL_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_OUT)
    print(f"Generado: {EXCEL_OUT}")


if __name__ == "__main__":
    main()
